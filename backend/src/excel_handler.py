from openpyxl import Workbook
from openpyxl import load_workbook
import xlwt
import xlrd


def table_to_xlsx(headers: list, rows: list, output_path: str):
    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(output_path)


def table_to_xls(headers: list, rows: list, output_path: str):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    for col_idx, header in enumerate(headers):
        ws.write(0, col_idx, header)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            ws.write(row_idx, col_idx, value)
    wb.save(output_path)


def read_excel(file_path: str) -> dict:
    headers, rows = [], []
    if file_path.lower().endswith(".xlsx"):
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            values = [str(c) if c is not None else "" for c in row]
            if row_idx == 0:
                headers = values
            else:
                rows.append(values)
    else:
        wb = xlrd.open_workbook(file_path)
        ws = wb.sheet_by_index(0)
        for row_idx in range(ws.nrows):
            values = [str(ws.cell_value(row_idx, c)) for c in range(ws.ncols)]
            if row_idx == 0:
                headers = values
            else:
                rows.append(values)
    max_c = max(len(headers), max((len(r) for r in rows), default=0))
    if len(headers) < max_c:
        headers += [""] * (max_c - len(headers))
    rows = [r + [""] * (max_c - len(r)) for r in rows]
    return {"headers": headers, "rows": rows}
